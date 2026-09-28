import csv
import os
import re
from xml.sax.saxutils import escape as xml_escape

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.schemas.analytics_reports import (
    ExportFormat, IndividualPerformanceReport, CompanyWidePerformanceReport, ReportScope,
)

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "generated_reports")
os.makedirs(OUTPUT_DIR, exist_ok=True)

EXTENSIONS = {ExportFormat.PDF: "pdf", ExportFormat.EXCEL: "xlsx", ExportFormat.CSV: "csv"}
MEDIA_TYPES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}


def _p(value) -> str:
    """Text for a ReportLab Paragraph, which parses its own mini-HTML (a stray '<' breaks the PDF)."""
    return xml_escape(str(value if value is not None else ""))


def _safe(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "_", str(value))[:40]


def _cell(value):
    """Spreadsheet-safe value: text starting with =,+,-,@ would run as a formula in Excel."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


class ReportExportService:
    @classmethod
    def export(cls, report, scope: ReportScope, fmt: ExportFormat, job_id: str) -> str:
        """Writes the report in the requested format and returns the file name inside OUTPUT_DIR."""
        ext = EXTENSIONS[fmt]
        if scope == ReportScope.INDIVIDUAL:
            name = f"Employee_Performance_{_safe(report.employee_id)}_{job_id[:8]}.{ext}"
            writer = {ExportFormat.PDF: cls.generate_individual_pdf,
                      ExportFormat.EXCEL: cls.generate_individual_excel,
                      ExportFormat.CSV: cls.generate_individual_csv}[fmt]
        else:
            name = f"Company_Report_{job_id[:8]}.{ext}"
            writer = {ExportFormat.PDF: cls.generate_company_pdf,
                      ExportFormat.EXCEL: cls.generate_company_multi_sheet_excel,
                      ExportFormat.CSV: cls.generate_company_csv}[fmt]
        writer(report, job_id, os.path.join(OUTPUT_DIR, name))
        return name

    # ------------------------------------------------------------------ rows shared by Excel / CSV / PDF
    @staticmethod
    def _individual_rows(report: IndividualPerformanceReport):
        sc = report.score_card
        return [
            ["Employee", f"{report.employee_name} ({report.employee_id})"],
            ["Department", report.department],
            ["Designation", report.designation],
            ["Period", f"{report.start_date} to {report.end_date}"],
            ["Overall score", sc.overall_score],
            ["Grade", sc.grade],
            ["Attendance score (of 30)", sc.attendance_score],
            ["Client hours score (of 30)", sc.productivity_hours_score],
            ["Task delivery score (of 30)", sc.task_completion_score],
            ["Blocker mitigation score (of 10)", sc.blocker_mitigation_score],
            ["Present days", report.total_present_days],
            ["Late days", report.total_late_days],
            ["Absent days", report.total_absent_days],
            ["Leave days", report.total_leave_days],
            ["Client hours logged", report.total_client_hours_logged],
            ["Average daily hours", report.average_daily_hours],
            ["Tasks total", report.tasks_total],
            ["Tasks completed", report.tasks_completed],
            ["Tasks in progress", report.tasks_in_progress],
            ["Tasks blocked", report.tasks_blocked],
            ["Task completion rate (%)", report.task_completion_rate],
        ]

    @staticmethod
    def _company_summary_rows(report: CompanyWidePerformanceReport):
        return [
            ["Report window", report.date_range_label],
            ["Total active staff", report.total_active_employees],
            ["Total billed client hours", report.total_hours_billed_clients],
            ["Company efficiency rate (%)", report.company_wide_efficiency_rate],
            ["Company average score (of 100)", report.company_avg_performance_score],
        ]

    DEPT_HEADER = ["Department", "Headcount", "Total Logged Hours", "Avg Hours / Employee", "Score Average",
                   "Tasks Completed", "Active Blockers"]

    @staticmethod
    def _dept_row(d):
        return [d.department, d.headcount, d.total_hours_logged, d.avg_hours_per_employee,
                d.average_performance_score, d.tasks_completed_count, d.active_blockers_count]

    # ------------------------------------------------------------------ individual
    @classmethod
    def generate_individual_excel(cls, report: IndividualPerformanceReport, job_id: str, file_path: str) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Scorecard"
        ws.append(["REXERA HRMS - EMPLOYEE PERFORMANCE SCORECARD"])
        ws["A1"].font = Font(bold=True, size=13)
        for row in cls._individual_rows(report):
            ws.append([_cell(v) for v in row])
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 40
        daily = wb.create_sheet("Daily Hours")
        daily.append(["Date", "Hours"])
        data = report.daily_hours_chart.datasets[0].data if report.daily_hours_chart.datasets else []
        for label, value in zip(report.daily_hours_chart.labels, data):
            daily.append([_cell(label), value])
        wb.save(file_path)
        return file_path

    @classmethod
    def generate_individual_csv(cls, report: IndividualPerformanceReport, job_id: str, file_path: str) -> str:
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["Metric", "Value"])
            for row in cls._individual_rows(report):
                w.writerow([_cell(v) for v in row])
        return file_path

    # ------------------------------------------------------------------ company
    @classmethod
    def generate_company_csv(cls, report: CompanyWidePerformanceReport, job_id: str, file_path: str) -> str:
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["Metric", "Value"])
            for row in cls._company_summary_rows(report):
                w.writerow([_cell(v) for v in row])
            w.writerow([])
            w.writerow(cls.DEPT_HEADER)
            for d in report.department_metrics:
                w.writerow([_cell(v) for v in cls._dept_row(d)])
        return file_path

    @classmethod
    def generate_company_pdf(cls, report: CompanyWidePerformanceReport, job_id: str, file_path: str) -> str:
        doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        elements = [
            Paragraph("REXERA TECHNOLOGIES", styles["Title"]),
            Paragraph(f"Company Performance Summary &bull; {_p(report.date_range_label)}", styles["Normal"]),
            Spacer(1, 12),
        ]
        summary = Table([[_p(k), _p(v)] for k, v in cls._company_summary_rows(report)], colWidths=[250, 290])
        summary.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#f8f9fa")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#dee2e6")),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        elements += [summary, Spacer(1, 15), Paragraph("<b>Department breakdown</b>", styles["Heading3"])]
        rows = [["Department", "Staff", "Hours", "Avg hrs", "Score"]] + [
            [Paragraph(_p(d.department), styles["Normal"]), d.headcount, d.total_hours_logged,
             d.avg_hours_per_employee, d.average_performance_score]
            for d in report.department_metrics
        ]
        dept = Table(rows, colWidths=[180, 70, 90, 90, 110])
        dept.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2b3a4a")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e0e0e0")),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(dept)
        doc.build(elements)
        return file_path

    @classmethod
    def generate_individual_pdf(cls, report: IndividualPerformanceReport, job_id: str, file_path: str = None) -> str:
        file_path = file_path or os.path.join(OUTPUT_DIR, f"Employee_Performance_{_safe(report.employee_id)}_{job_id[:8]}.pdf")
        doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1a2e3b")
        )
        subtitle_style = ParagraphStyle(
            "SubTitleStyle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#6c757d")
        )

        elements.append(Paragraph("REXERA TECHNOLOGIES", title_style))
        elements.append(Paragraph(f"Employee Performance &amp; Productivity Scorecard &bull; {_p(report.date_range_label)}", subtitle_style))
        elements.append(Spacer(1, 15))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#007bff"), spaceAfter=15))

        info_data = [
            [
                Paragraph(f"<b>Employee:</b> {_p(report.employee_name)} ({_p(report.employee_id)})", styles["Normal"]),
                Paragraph(f"<b>Overall Score:</b> <font color='#28a745'><b>{report.score_card.overall_score}/100</b></font>", styles["Normal"])
            ],
            [
                Paragraph(f"<b>Department:</b> {_p(report.department)}", styles["Normal"]),
                Paragraph(f"<b>Performance Grade:</b> <b>{_p(report.score_card.grade)}</b>", styles["Normal"])
            ],
            [
                Paragraph(f"<b>Designation:</b> {_p(report.designation)}", styles["Normal"]),
                Paragraph(f"<b>Period:</b> {_p(report.start_date)} to {_p(report.end_date)}", styles["Normal"])
            ]
        ]
        info_table = Table(info_data, colWidths=[270, 270])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8f9fa")),
            ('PADDING', (0, 0), (-1, -1), 6),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#dee2e6"))
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 15))

        elements.append(Paragraph("<b>Performance Pillar Breakdown</b>", styles["Heading3"]))
        elements.append(Spacer(1, 5))
        score_data = [
            ["Evaluation Pillar", "Weight / Max", "Score Earned", "Status"],
            ["Attendance & Punctuality", "30 Pts", f"{report.score_card.attendance_score} Pts", "Excellent" if report.score_card.attendance_score > 24 else "Normal"],
            ["Client Productive Hours", "30 Pts", f"{report.score_card.productivity_hours_score} Pts", f"{report.total_client_hours_logged} hrs logged"],
            ["Task Velocity & Delivery", "30 Pts", f"{report.score_card.task_completion_score} Pts", f"{report.task_completion_rate}% completed"],
            ["Blocker Mitigation", "10 Pts", f"{report.score_card.blocker_mitigation_score} Pts", f"{report.tasks_blocked} active blockers"]
        ]
        score_table = Table(score_data, colWidths=[180, 110, 120, 130])
        score_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2b3a4a")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e0e0e0")),
            ('PADDING', (0, 0), (-1, -1), 6)
        ]))
        elements.append(score_table)

        doc.build(elements)
        return file_path

    @classmethod
    def generate_company_multi_sheet_excel(
        cls,
        report: CompanyWidePerformanceReport,
        job_id: str,
        file_path: str = None,
    ) -> str:
        wb = openpyxl.Workbook()
        
        ws_sum = wb.active
        ws_sum.title = "Executive Summary"
        
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")

        ws_sum.append(["REXERA HRMS - COMPANY PERFORMANCE SUMMARY"])
        ws_sum.append(["Report Window:", report.date_range_label])
        ws_sum.append(["Total Active Staff:", report.total_active_employees])
        ws_sum.append(["Total Billed Client Hours:", report.total_hours_billed_clients])
        ws_sum.append(["Company Efficiency Rate:", f"{report.company_wide_efficiency_rate}%"])
        ws_sum.append(["Company Avg Score:", f"{report.company_avg_performance_score}/100"])

        ws_dept = wb.create_sheet(title="Department Breakdown")
        ws_dept.append(["Department", "Headcount", "Total Logged Hours", "Avg Hours / Employee", "Score Average"])
        for c in ws_dept[1]:
            c.fill = header_fill
            c.font = header_font

        for d in report.department_metrics:
            ws_dept.append([_cell(d.department), d.headcount, d.total_hours_logged, d.avg_hours_per_employee, d.average_performance_score])

        ws_block = wb.create_sheet(title="Red Zone Blockers")
        ws_block.append(["Client", "Project", "Task", "Assignee", "Blocker Reason", "Blocked Since"])
        for c in ws_block[1]:
            c.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            c.font = header_font

        for b in report.top_critical_bottlenecks:
            ws_block.append([_cell(b.get(k)) for k in ("client_name", "project_name", "task_title", "assigned_to", "reason", "blocked_at")])

        file_path = file_path or os.path.join(OUTPUT_DIR, f"Company_Report_{job_id[:8]}.xlsx")
        wb.save(file_path)
        return file_path
