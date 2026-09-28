import os
import re
import io
import uuid
import logging
from datetime import datetime
from typing import Dict, Any, List, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.database import get_collection
from app.schemas.leave import LeaveBalance
from app.services.employee_service import EmployeeService
from app.utils.validators import optional_ifsc, validate_mobile
from app.schemas.bulk_import import (
    ImportTargetEntity,
    ImportJobStatus,
    ImportJobSummary,
    TARGET_SCHEMA_REGISTRY
)

logger = logging.getLogger("rexera.bulk_import.engine")

ERROR_REPORT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "error_reports")


class ImportExecutionEngine:
    @staticmethod
    def validate_field_value(val: Any, f_type: str, is_required: bool) -> Tuple[bool, Any, Optional[str]]:
        val_str = str(val).strip() if val is not None else ""
        if not val_str:
            if is_required:
                return False, None, "Mandatory field is missing or empty."
            return True, None, None

        if f_type == "email":
            if not re.match(r"^[\w\.-]+@[\w\.-]+\.\w+$", val_str):
                return False, None, f"Invalid email format: '{val_str}'"
            return True, val_str.lower(), None

        elif f_type == "float":
            # Allow "₹1,50,000", "Rs. 5000", "INR 5000"; stripping every non-digit used to turn -5000 into 5000.
            clean_num = re.sub(r"(?i)(inr|rs\.?|₹|,|\s)", "", val_str)
            try:
                num = float(clean_num)
            except ValueError:
                return False, None, f"Invalid numeric value: '{val_str}'"
            if num < 0:
                return False, None, f"Amount cannot be negative: '{val_str}'"
            return True, num, None

        elif f_type == "date":
            # Excel date cells arrive as "2026-09-01 00:00:00" (read as text); typed dates are Indian dd/mm/yyyy.
            candidates = [val_str[:10]] if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", val_str) else [val_str]
            for candidate in candidates:
                for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"]:
                    try:
                        dt = datetime.strptime(candidate, fmt)
                        return True, dt.strftime("%Y-%m-%d"), None
                    except ValueError:
                        pass
            return False, None, f"Invalid date format: '{val_str}'. Expected YYYY-MM-DD or DD/MM/YYYY."

        elif f_type == "phone":
            ok, res = validate_mobile(val_str)
            return (True, res, None) if ok else (False, None, res)

        elif f_type == "ifsc":
            try:
                return True, optional_ifsc(val_str) or None, None
            except ValueError as e:
                return False, None, str(e)

        return True, val_str, None

    @classmethod
    def generate_error_report_excel(
        cls,
        failed_records: List[Dict[str, Any]],
        job_id: str
    ) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Import_Errors"

        if not failed_records:
            return ""

        raw_headers = [k for k in failed_records[0].keys() if k not in ["_row_index", "_error_reason"]]
        headers = ["Row #", "Failure Reason"] + raw_headers

        header_fill = PatternFill(start_color="DC3545", end_color="DC3545", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in failed_records:
            row_data = [
                row.get("_row_index"),
                row.get("_error_reason")
            ] + [row.get(h, "") for h in raw_headers]
            ws.append(row_data)

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        os.makedirs(ERROR_REPORT_DIR, exist_ok=True)
        file_path = os.path.join(ERROR_REPORT_DIR, f"error_report_{job_id}.xlsx")
        wb.save(file_path)

        return f"/api/bulk-import/error-reports/{job_id}"

    @classmethod
    async def _complete_new_employee(cls, doc: Dict[str, Any], col) -> Dict[str, Any]:
        """Fills the fields every employee record needs so imported rows show up in the directory and payroll."""
        code = doc.get("employee_code")
        if not code:
            code = await EmployeeService.generate_next_employee_code()
        full = {
            "employee_code": code, "mobile_number": "", "reporting_manager": "", "date_of_joining": "",
            "base_salary": 0.0, "hra": 0.0, "conveyance_allowance": 0.0, "special_allowance": 0.0,
            "professional_tax": 200.0, "pf_opted": True, "bank_name": "", "account_no": "", "ifsc_code": "",
            "employee_status": "Active", "joining_status": "Joined",
        }
        full.update({k: v for k, v in doc.items() if v is not None})
        full.update(cls._salary_totals(full))
        return full

    @staticmethod
    def _salary_totals(emp: Dict[str, Any]) -> Dict[str, float]:
        num = lambda k, d=0.0: float(emp.get(k) if emp.get(k) is not None else d)
        gross = num("base_salary") + num("hra") + num("conveyance_allowance") + num("special_allowance")
        pf = num("base_salary") * 0.12 if emp.get("pf_opted", True) else 0.0
        return {"gross_salary": gross, "estimated_net_salary": max(0.0, gross - pf - num("professional_tax", 200.0))}

    @classmethod
    async def _refresh_employee_salary(cls, emp_id: str, col) -> None:
        """An import that changes salary fields must also update the totals and the payroll structure."""
        emp = await col.find_one({"_id": emp_id})
        if emp:
            await col.update_one({"_id": emp_id}, {"$set": cls._salary_totals(emp)})
            await EmployeeService.sync_salary_structure(emp_id)

    @classmethod
    async def execute_upsert_batch(
        cls,
        job_id: str,
        records: List[Dict[str, Any]],
        mappings: Dict[str, str],
        target_entity: ImportTargetEntity,
        unique_key_field: str,
        is_dry_run: bool,
        user_email: str
    ):
        job_col = get_collection("import_jobs")
        target_col_name = "employees" if target_entity == ImportTargetEntity.EMPLOYEES else "leave_balances"
        target_col = get_collection(target_col_name)

        schema_meta = TARGET_SCHEMA_REGISTRY.get(target_entity, {}).get("fields", {})
        total_rows = len(records)
        balance_year = datetime.utcnow().year
        
        inserted = 0
        updated = 0
        failed_rows = []

        await job_col.update_one(
            {"job_id": job_id},
            {"$set": {"status": ImportJobStatus.PROCESSING.value, "total_rows": total_rows}}
        )

        for idx, row in enumerate(records, start=1):
            row_errors = []
            clean_doc = {}

            for file_header, db_field in mappings.items():
                if not db_field or db_field == "IGNORE":
                    continue

                field_info = schema_meta.get(db_field, {"required": False, "type": "str"})
                raw_val = row.get(file_header)

                is_valid, parsed_val, err_msg = cls.validate_field_value(
                    raw_val, field_info.get("type", "str"), field_info.get("required", False)
                )
                if not is_valid:
                    row_errors.append(f"[{file_header}]: {err_msg}")
                else:
                    clean_doc[db_field] = parsed_val

            key_field = unique_key_field
            if target_entity == ImportTargetEntity.EMPLOYEES:
                # Employees match on their code when given, otherwise on email.
                key_field = "employee_code" if clean_doc.get("employee_code") else "email"
            elif target_entity == ImportTargetEntity.LEAVE_BALANCES:
                key_field = "employee_id"

            if key_field not in clean_doc or not clean_doc[key_field]:
                row_errors.append(f"Unique Key '{key_field}' is missing.")

            filter_q: Dict[str, Any] = {key_field: clean_doc.get(key_field)}
            if target_entity == ImportTargetEntity.LEAVE_BALANCES:
                # Balances are per employee *and year*; without the year the app never found imported rows.
                filter_q["year"] = balance_year
            if target_entity == ImportTargetEntity.EMPLOYEES and not row_errors and clean_doc.get("email") and key_field == "employee_code":
                other = await target_col.find_one({
                    "email": {"$regex": f"^{re.escape(clean_doc['email'])}$", "$options": "i"},
                    "employee_code": {"$ne": clean_doc["employee_code"]},
                })
                if other:
                    row_errors.append(f"Email {clean_doc['email']} already belongs to {other.get('employee_code')}.")

            if row_errors:
                failed_entry = dict(row)
                failed_entry["_row_index"] = idx
                failed_entry["_error_reason"] = " | ".join(row_errors)
                failed_rows.append(failed_entry)
                continue

            if not is_dry_run:
                existing = await target_col.find_one(filter_q)
                now_utc = datetime.utcnow()

                if existing:
                    clean_doc["updated_at"] = now_utc
                    await target_col.update_one({"_id": existing["_id"]}, {"$set": clean_doc})
                    if target_entity == ImportTargetEntity.EMPLOYEES:
                        await cls._refresh_employee_salary(existing["_id"], target_col)
                    updated += 1
                else:
                    if target_entity == ImportTargetEntity.EMPLOYEES:
                        clean_doc = await cls._complete_new_employee(clean_doc, target_col)
                    elif target_entity == ImportTargetEntity.LEAVE_BALANCES:
                        clean_doc = {**LeaveBalance(employee_id=clean_doc["employee_id"], year=balance_year).model_dump(exclude={"id"}),
                                     **clean_doc}
                    clean_doc["created_at"] = now_utc
                    clean_doc["updated_at"] = now_utc
                    await target_col.insert_one(clean_doc)
                    inserted += 1
            else:
                existing = await target_col.find_one(filter_q)
                if existing:
                    updated += 1
                else:
                    inserted += 1

            if idx % 50 == 0 or idx == total_rows:
                progress = int((idx / total_rows) * 100)
                await job_col.update_one(
                    {"job_id": job_id},
                    {"$set": {"progress_percentage": progress}}
                )

        error_url = None
        if failed_rows:
            error_url = cls.generate_error_report_excel(failed_rows, job_id)

        now = datetime.utcnow()
        await job_col.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": ImportJobStatus.COMPLETED.value,
                    "progress_percentage": 100,
                    "inserted_count": inserted,
                    "updated_count": updated,
                    "failed_count": len(failed_rows),
                    "error_report_file_url": error_url,
                    "completed_at": now
                }
            }
        )

        return {
            "job_id": job_id,
            "total": total_rows,
            "inserted": inserted,
            "updated": updated,
            "failed": len(failed_rows),
            "error_url": error_url
        }
